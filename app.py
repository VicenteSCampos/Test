import sys
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import datetime

import anthropic
from dotenv import load_dotenv

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QDialog,
    QVBoxLayout, QHBoxLayout, QGridLayout,
    QPushButton, QLabel, QLineEdit, QComboBox,
    QRadioButton, QButtonGroup, QProgressBar,
    QTextEdit, QTableWidget, QTableWidgetItem,
    QTabWidget, QScrollArea, QFileDialog,
    QMessageBox, QHeaderView, QAbstractItemView,
    QFrame, QSizePolicy,
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QColor, QFont, QTextCursor

_base = os.path.dirname(os.path.abspath(__file__))
load_dotenv(dotenv_path=os.path.join(_base, ".env"))

logging.basicConfig(
    filename=os.path.join(_base, "error.log"),
    level=logging.DEBUG,
    format="%(asctime)s %(levelname)s: %(message)s",
)

DARK_STYLESHEET = """
QMainWindow, QDialog { background-color: #1e1e1e; }
QWidget { background-color: #1e1e1e; color: #ffffff; font-size: 10pt; }
QTabWidget::pane { border: none; background-color: #1e1e1e; }
QTabBar::tab {
    background-color: #2b2b2b; color: #ffffff;
    padding: 6px 16px; border: none; font-size: 10pt;
}
QTabBar::tab:selected { background-color: #1a4a7a; }
QTabBar::tab:hover:!selected { background-color: #2d5a8a; }
QPushButton {
    background-color: #1f538d; color: #ffffff;
    border: none; padding: 5px 10px; border-radius: 4px; font-size: 10pt;
}
QPushButton:hover { background-color: #2d5a8a; }
QPushButton:pressed { background-color: #163d6b; }
QPushButton:disabled { background-color: #3a3a3a; color: #666666; }
QPushButton[class="gray"] { background-color: #4a4a4a; }
QPushButton[class="gray"]:hover { background-color: #5a5a5a; }
QPushButton[class="gray"]:pressed { background-color: #3a3a3a; }
QPushButton[class="danger"] { background-color: #7a1a1a; }
QPushButton[class="danger"]:hover { background-color: #a02020; }
QLineEdit, QComboBox {
    background-color: #2b2b2b; color: #ffffff;
    border: 1px solid #4a4a4a; padding: 4px 8px;
    border-radius: 4px; font-size: 10pt;
}
QComboBox::drop-down { border: none; width: 20px; background-color: #3a3a3a; }
QComboBox QAbstractItemView {
    background-color: #2b2b2b; color: white;
    selection-background-color: #1a4a7a;
}
QTextEdit {
    background-color: #1a1a1a; color: #d4d4d4;
    border: 1px solid #3a3a3a; border-radius: 4px;
    font-family: Consolas, monospace; font-size: 10pt;
}
QTableWidget {
    background-color: #2b2b2b; color: #ffffff;
    gridline-color: #3a3a3a; border: none;
}
QTableWidget::item { padding: 4px; border: none; }
QTableWidget::item:selected { background-color: #1a5276; color: white; }
QHeaderView::section {
    background-color: #1a4a7a; color: #ffffff;
    font-weight: bold; border: none; padding: 4px 6px; font-size: 10pt;
}
QProgressBar {
    background-color: #2b2b2b; color: white;
    border: 1px solid #4a4a4a; border-radius: 4px; text-align: center;
}
QProgressBar::chunk { background-color: #1f538d; border-radius: 3px; }
QScrollArea { border: none; background-color: transparent; }
QScrollBar:vertical {
    background-color: #2b2b2b; width: 12px; border: none;
}
QScrollBar::handle:vertical {
    background-color: #4a4a4a; border-radius: 6px; min-height: 20px;
}
QScrollBar::handle:vertical:hover { background-color: #5a5a5a; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
QRadioButton { color: white; spacing: 6px; font-size: 10pt; background: transparent; }
QLabel { color: white; background: transparent; }
"""


@dataclass
class QueueItem:
    audio_paths: list = field(default_factory=list)
    asignatura: str = ""
    fecha: str = ""
    pptx_path: str = ""
    nombre: str = ""
    estado: str = "pendiente"  # pendiente / procesando / listo / error


def extract_pptx(filepath: str) -> str:
    from pptx import Presentation
    prs = Presentation(filepath)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            slides.append(f"[Slide {i}]: {' | '.join(texts)}")
    return "\n".join(slides)


def scan_vault_folders(vault: str) -> list:
    if not os.path.isdir(vault):
        return []
    return sorted([
        f for f in os.listdir(vault)
        if os.path.isdir(os.path.join(vault, f)) and not f.startswith(".")
    ])


def _strip_emoji(text: str) -> str:
    return re.sub(r"[^\x00-\x7F\u00C0-\u024F\u0400-\u04FF]", "", text).strip()


def _pdf_safe(text: str) -> str:
    return text.replace("\u00a0", " ")


def save_pdf(content: str, filepath: str):
    from fpdf import FPDF

    class PDF(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Arial", "I", 8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 10, f"Pagina {self.page_no()}", align="C")

    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_margins(20, 20, 20)
    pdf.add_font("Arial", fname="C:\\Windows\\Fonts\\arial.ttf")
    pdf.add_font("Arial", style="B", fname="C:\\Windows\\Fonts\\arialbd.ttf")
    pdf.add_font("Arial", style="I", fname="C:\\Windows\\Fonts\\ariali.ttf")

    lines = content.split("\n")
    start, in_front = 0, False
    for i, line in enumerate(lines):
        if i == 0 and line.strip() == "---":
            in_front = True
            continue
        if in_front and line.strip() == "---":
            start = i + 1
            break

    lm = pdf.l_margin
    pw = pdf.w - pdf.l_margin - pdf.r_margin

    in_mermaid = False
    for line in lines[start:]:
        s = _pdf_safe(_strip_emoji(line.rstrip())).strip()

        if s.startswith("```mermaid"):
            in_mermaid = True
            continue
        if in_mermaid:
            if s == "```":
                in_mermaid = False
            continue

        pdf.set_x(lm)

        if not s:
            pdf.ln(2)
        elif s.startswith("## "):
            pdf.set_font("Arial", "B", 14)
            pdf.set_text_color(30, 120, 200)
            pdf.ln(4)
            pdf.multi_cell(pw, 8, s[3:])
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        elif s.startswith("### "):
            pdf.set_font("Arial", "B", 11)
            pdf.ln(3)
            pdf.multi_cell(pw, 7, s[4:])
            pdf.ln(1)
        elif s.startswith("> "):
            pdf.set_font("Arial", "I", 10)
            pdf.set_fill_color(240, 240, 240)
            indent = 8
            pdf.set_x(lm + indent)
            pdf.multi_cell(pw - indent, 6, s[2:], fill=True)
        elif re.match(r"^[-*] ", s):
            pdf.set_font("Arial", "", 10)
            indent = 6
            pdf.set_x(lm + indent)
            pdf.multi_cell(pw - indent, 6, f"• {s[2:]}")
        elif re.match(r"^\d+\.", s):
            pdf.set_font("Arial", "", 10)
            indent = 6
            pdf.set_x(lm + indent)
            pdf.multi_cell(pw - indent, 6, s)
        elif s == "---":
            pdf.ln(2)
            pdf.set_draw_color(180, 180, 180)
            pdf.line(lm, pdf.get_y(), lm + pw, pdf.get_y())
            pdf.ln(2)
        else:
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
            pdf.set_font("Arial", "", 10)
            pdf.multi_cell(pw, 6, clean)

    pdf.output(filepath)


class _Worker(QThread):
    log_signal = pyqtSignal(str)
    progress_signal = pyqtSignal(float, str)
    refresh_signal = pyqtSignal()
    done_signal = pyqtSignal()

    def __init__(self, app: "ClaseObsidianApp"):
        super().__init__()
        self._app = app

    def _log(self, msg: str):
        self.log_signal.emit(msg)

    def _set_progress(self, value: float, label: str = ""):
        self.progress_signal.emit(value, label)

    def run(self):
        app = self._app
        pending = [i for i, item in enumerate(app.queue) if item.estado == "pendiente"]
        total = len(pending)
        success = 0
        try:
            for step, idx in enumerate(pending):
                item = app.queue[idx]
                item.estado = "procesando"
                self.refresh_signal.emit()

                try:
                    n = len(item.audio_paths)
                    label = os.path.basename(item.audio_paths[0]) if n == 1 else f"{n} archivos"
                    self._log(f"\n[{step + 1}/{total}] 🎵 {label}")

                    transcription = self._transcribe(item.audio_paths)

                    pptx_text = ""
                    if item.pptx_path:
                        self._log("📊 Extrayendo texto del PowerPoint...")
                        pptx_text = extract_pptx(item.pptx_path)
                        self._log(f"✅ PPT procesado — {len(pptx_text.split())} palabras")

                    is_obsidian_md = (app._dest == "obsidian" and app._fmt == "md")
                    content = self._call_claude(
                        transcription, item.asignatura, item.fecha,
                        pptx_text=pptx_text, include_mermaid=is_obsidian_md,
                    )
                    full = content + f"\n\n---\n\n## 📄 Transcripción\n\n{transcription}"
                    saved_path = app._save(full, item.asignatura, item.fecha, item.nombre)
                    self._log(f"✅ Guardado: {os.path.basename(saved_path)}")

                    item.estado = "listo"
                    success += 1
                    self._set_progress((step + 1) / total, f"[{step + 1}/{total}] Completado")

                except Exception as e:
                    logging.exception(f"Error procesando item {idx}")
                    self._log(f"❌ Error: {type(e).__name__}: {e}")
                    item.estado = "error"

                self.refresh_signal.emit()

            self._set_progress(1.0, f"¡Completado! {success}/{total} exitosos")
        finally:
            self.done_signal.emit()

    def _transcribe(self, audio_paths: list) -> str:
        app = self._app
        from faster_whisper import WhisperModel

        if app._whisper_model is None:
            self._log("⏳ Cargando modelo Whisper...")
            self._set_progress(0.05, "Cargando modelo...")
            app._whisper_model = WhisperModel(app._model, device="cuda", compute_type="int8_float16")

        model = app._whisper_model
        all_parts = []

        for file_idx, audio_path in enumerate(audio_paths):
            prefix = f"[{file_idx + 1}/{len(audio_paths)}] " if len(audio_paths) > 1 else ""
            self._log(f"🎙️ {prefix}Transcribiendo {os.path.basename(audio_path)}...")
            self._set_progress(0.1, f"{prefix}Transcribiendo audio...")

            segments, info = model.transcribe(
                audio_path,
                language="es",
                beam_size=5,
                vad_filter=True,
                vad_parameters={"min_silence_duration_ms": 500},
                condition_on_previous_text=False,
                no_speech_threshold=0.6,
                compression_ratio_threshold=2.4,
                initial_prompt="Transcripción de clase universitaria en español.",
            )
            parts = []
            for seg in segments:
                parts.append(seg.text.strip())
                if info.duration > 0:
                    file_progress = file_idx / len(audio_paths)
                    seg_progress = (seg.end / info.duration) / len(audio_paths)
                    pct = 0.1 + (file_progress + seg_progress) * 0.5
                    self._set_progress(
                        min(pct, 0.6),
                        f"{prefix}Transcribiendo... "
                        f"{int(seg.end / 60)}:{int(seg.end % 60):02d} / "
                        f"{int(info.duration / 60)}:{int(info.duration % 60):02d}",
                    )
            all_parts.extend(parts)
            self._log(f"✅ {prefix}Listo — {len(parts)} segmentos")

        transcription = " ".join(all_parts)
        self._log(f"✅ Transcripción total — {len(transcription.split())} palabras")
        self._set_progress(0.65, "Procesando con Claude AI...")
        return transcription

    def _call_claude(self, transcription: str, asignatura: str, fecha: str,
                     pptx_text: str = "", include_mermaid: bool = False) -> str:
        self._log("🤖 Enviando a Claude AI para estructurar el documento...")

        api_key = os.getenv("ANTHROPIC_API_KEY")
        if not api_key:
            raise ValueError("No se encontró ANTHROPIC_API_KEY en el archivo .env")

        client = anthropic.Anthropic(api_key=api_key)

        ppt_block = f"\n\nCONTENIDO DEL POWERPOINT:\n{pptx_text}" if pptx_text else ""

        mermaid_block = (
            "\n## 🗺️ Mapa Conceptual\n\n"
            "```mermaid\n"
            "graph TD\n"
            "    [Genera el diagrama aquí con máximo 10 nodos representando los conceptos principales y sus relaciones]\n"
            "```\n"
        ) if include_mermaid else ""

        tag = asignatura.lower().replace(" ", "-")

        prompt = (
            f"Eres un asistente académico experto. Analiza la siguiente información de una clase universitaria "
            f"y genera un documento en Markdown con EXACTAMENTE esta estructura:\n\n"
            f"---\n"
            f"asignatura: {asignatura}\n"
            f"fecha: {fecha}\n"
            f"tema: [extrae el tema principal en máximo 5 palabras]\n"
            f"tags: [{tag}]\n"
            f"---\n\n"
            f"## 📋 Ficha Técnica\n"
            f"- **Fecha:** {fecha}\n"
            f"- **Asignatura:** {asignatura}\n"
            f"- **Tema Principal:** [extrae el tema central]\n"
            f"- **Objetivo de la Clase:** [extrae o infiere el objetivo principal]\n\n"
            f"## 📝 Resumen Ejecutivo\n"
            f"[Máximo 5 párrafos. Explica el contexto y el núcleo de la clase de forma comprensible "
            f"para alguien que no asistió. Si hay PPT, intégralo con la transcripción.]\n\n"
            f"## 🧠 Desarrollo Lógico de la Clase\n"
            f"[Describe en orden cronológico cómo se desarrolló la clase. Usa una lista numerada donde cada punto "
            f"represente una etapa o transición temática. Muestra cómo cada idea o bloque temático se conecta con "
            f"el siguiente, formando una progresión lógica. Ejemplo de formato:]\n\n"
            f"1. **Introducción / Contexto:** [cómo abrió la clase]\n"
            f"2. **Desarrollo central:** [los bloques temáticos principales en orden]\n"
            f"3. **Conexiones y transiciones:** [cómo se encadenaron los temas]\n"
            f"4. **Cierre / Conclusión:** [cómo terminó o qué se dejó abierto]\n\n"
            f"## 🔑 Conceptos Clave y Definiciones\n"
            f"[Para cada concepto importante usa este formato exacto:]\n\n"
            f"**Término:** Definición clara y concisa.\n"
            f"> Ejemplo: [si el profesor o el PPT dieron un ejemplo]\n\n"
            f"## 📊 Marco Lógico\n"
            f"[Si la clase presenta un modelo de negocio, estrategia empresarial o política pública:]\n"
            f"- **Objetivo:**\n"
            f"- **Medios:**\n"
            f"- **Resultados:**\n"
            f"- **Supuestos:**\n"
            f"[Si NO aplica escribe únicamente: \"No aplica para esta clase.\"]\n"
            f"{mermaid_block}\n"
            f"## ❓ Preguntas de Examen\n"
            f"[Entre 5 y 8 preguntas desafiantes. Usa EXACTAMENTE este formato para cada una:]\n\n"
            f"> [!question]- [escribe aquí la pregunta]\n"
            f"> [escribe aquí la respuesta detallada]\n\n"
            f"---\n"
            f"TRANSCRIPCIÓN:\n{transcription}"
            f"{ppt_block}"
        )

        msg = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}],
        )
        self._set_progress(0.9, "Guardando documento...")
        return msg.content[0].text


class EditQueueItemDialog(QDialog):
    def __init__(self, parent, item: QueueItem = None, vault_folders: list = None):
        super().__init__(parent)
        self.setWindowTitle("Configurar elemento")
        self.setFixedSize(520, 580)
        self.result_item: QueueItem | None = None
        self._audio_paths: list = list(item.audio_paths) if item else []
        self._pptx_path = item.pptx_path if item else ""
        self._vault_folders = vault_folders or []
        self._setup_ui(item)

    def _setup_ui(self, item: QueueItem = None):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        # Scroll area for form
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll_content = QWidget()
        form = QVBoxLayout(scroll_content)
        form.setContentsMargins(20, 16, 20, 16)
        form.setSpacing(4)
        scroll.setWidget(scroll_content)
        outer.addWidget(scroll, 1)

        # Audio section header
        audio_header = QHBoxLayout()
        audio_lbl = QLabel("Archivos de Audio")
        audio_lbl.setStyleSheet("font-weight: bold;")
        btn_add = QPushButton("+ Agregar audio")
        btn_add.setFixedWidth(120)
        btn_add.clicked.connect(self._add_audio)
        audio_header.addWidget(audio_lbl)
        audio_header.addStretch()
        audio_header.addWidget(btn_add)
        form.addLayout(audio_header)
        form.addSpacing(4)

        # Audio list container (scrollable internally via a fixed-height frame)
        audio_frame = QFrame()
        audio_frame.setStyleSheet("QFrame { background-color: #1a1a1a; border-radius: 4px; }")
        audio_frame.setFixedHeight(130)
        audio_outer = QVBoxLayout(audio_frame)
        audio_outer.setContentsMargins(0, 0, 0, 0)

        audio_scroll = QScrollArea()
        audio_scroll.setWidgetResizable(True)
        audio_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        audio_scroll.setStyleSheet("QScrollArea { background: transparent; border: none; }")
        self._audio_container = QWidget()
        self._audio_container.setStyleSheet("background: transparent;")
        self._audio_layout = QVBoxLayout(self._audio_container)
        self._audio_layout.setContentsMargins(8, 8, 8, 8)
        self._audio_layout.setSpacing(4)
        self._audio_layout.addStretch()
        audio_scroll.setWidget(self._audio_container)
        audio_outer.addWidget(audio_scroll)
        form.addWidget(audio_frame)
        form.addSpacing(10)

        self._render_audio_list()

        # PPT section
        ppt_lbl = QLabel("PowerPoint (opcional)")
        ppt_lbl.setStyleSheet("font-weight: bold;")
        form.addWidget(ppt_lbl)
        form.addSpacing(4)

        ppt_row = QHBoxLayout()
        self._pptx_label = QLabel(
            os.path.basename(self._pptx_path) if self._pptx_path else "Ningún archivo seleccionado"
        )
        if not self._pptx_path:
            self._pptx_label.setStyleSheet("color: gray;")
        self._pptx_label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        btn_pptx = QPushButton("Seleccionar")
        btn_pptx.setFixedWidth(100)
        btn_pptx.clicked.connect(self._select_pptx)
        btn_clear = QPushButton("✕")
        btn_clear.setFixedWidth(30)
        btn_clear.setProperty("class", "gray")
        btn_clear.clicked.connect(self._clear_pptx)
        ppt_row.addWidget(self._pptx_label)
        ppt_row.addWidget(btn_pptx)
        ppt_row.addWidget(btn_clear)
        form.addLayout(ppt_row)
        form.addSpacing(10)

        # Asignatura
        asig_lbl = QLabel("Asignatura")
        asig_lbl.setStyleSheet("font-weight: bold;")
        form.addWidget(asig_lbl)
        form.addSpacing(4)
        self._asig_combo = QComboBox()
        self._asig_combo.setEditable(True)
        self._asig_combo.addItems(self._vault_folders)
        if item and item.asignatura:
            self._asig_combo.setCurrentText(item.asignatura)
        form.addWidget(self._asig_combo)
        form.addSpacing(10)

        # Fecha
        fecha_lbl = QLabel("Fecha")
        fecha_lbl.setStyleSheet("font-weight: bold;")
        form.addWidget(fecha_lbl)
        form.addSpacing(4)
        self._fecha_entry = QLineEdit()
        self._fecha_entry.setText(item.fecha if item else datetime.now().strftime("%Y-%m-%d"))
        form.addWidget(self._fecha_entry)
        form.addSpacing(10)

        # Nombre del archivo
        nombre_lbl = QLabel("Nombre del archivo (opcional)")
        nombre_lbl.setStyleSheet("font-weight: bold;")
        form.addWidget(nombre_lbl)
        hint_lbl = QLabel("Si se deja vacío se usará: fecha - asignatura")
        hint_lbl.setStyleSheet("color: gray; font-size: 9pt;")
        form.addWidget(hint_lbl)
        form.addSpacing(4)
        self._nombre_entry = QLineEdit()
        self._nombre_entry.setPlaceholderText("ej: Clase 1 - Introducción")
        self._nombre_entry.setText(item.nombre if item else "")
        form.addWidget(self._nombre_entry)
        form.addStretch()

        # Buttons — always visible at bottom
        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet("color: #3a3a3a;")
        outer.addWidget(sep)

        btn_row = QHBoxLayout()
        btn_row.setContentsMargins(20, 10, 20, 14)
        btn_row.addStretch()
        btn_cancel = QPushButton("Cancelar")
        btn_cancel.setProperty("class", "gray")
        btn_cancel.setFixedWidth(100)
        btn_cancel.clicked.connect(self.reject)
        btn_ok = QPushButton("Aceptar")
        btn_ok.setFixedWidth(100)
        btn_ok.clicked.connect(self._accept)
        btn_row.addWidget(btn_cancel)
        btn_row.addSpacing(8)
        btn_row.addWidget(btn_ok)
        outer.addLayout(btn_row)

    def _render_audio_list(self):
        while self._audio_layout.count():
            child = self._audio_layout.takeAt(0)
            if child.widget():
                child.widget().deleteLater()

        if not self._audio_paths:
            lbl = QLabel("Sin archivos — pulsa '+ Agregar audio'")
            lbl.setStyleSheet("color: gray; padding: 6px;")
            self._audio_layout.addWidget(lbl)
            self._audio_layout.addStretch()
            return

        for i, path in enumerate(self._audio_paths):
            row = QWidget()
            row.setStyleSheet("background: transparent;")
            row_layout = QHBoxLayout(row)
            row_layout.setContentsMargins(0, 0, 0, 0)
            row_layout.setSpacing(4)

            btn_up = QPushButton("↑")
            btn_up.setFixedSize(26, 26)
            btn_up.setProperty("class", "gray")
            btn_up.clicked.connect(lambda _, idx=i: self._move_audio(idx, -1))

            btn_down = QPushButton("↓")
            btn_down.setFixedSize(26, 26)
            btn_down.setProperty("class", "gray")
            btn_down.clicked.connect(lambda _, idx=i: self._move_audio(idx, 1))

            name = os.path.basename(path)
            if len(name) > 36:
                name = name[:33] + "..."
            name_lbl = QLabel(f"{i + 1}. {name}")
            name_lbl.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)

            btn_rm = QPushButton("✕")
            btn_rm.setFixedSize(26, 26)
            btn_rm.setProperty("class", "danger")
            btn_rm.clicked.connect(lambda _, idx=i: self._remove_audio(idx))

            row_layout.addWidget(btn_up)
            row_layout.addWidget(btn_down)
            row_layout.addWidget(name_lbl)
            row_layout.addWidget(btn_rm)
            self._audio_layout.addWidget(row)

        self._audio_layout.addStretch()

    def _add_audio(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Seleccionar archivos de audio", "",
            "Audio (*.mp3 *.wav *.m4a *.ogg *.flac *.mp4 *.mkv)",
        )
        for p in paths:
            if p not in self._audio_paths:
                self._audio_paths.append(p)
        self._render_audio_list()

    def _remove_audio(self, idx: int):
        self._audio_paths.pop(idx)
        self._render_audio_list()

    def _move_audio(self, idx: int, direction: int):
        new_idx = idx + direction
        if 0 <= new_idx < len(self._audio_paths):
            self._audio_paths[idx], self._audio_paths[new_idx] = (
                self._audio_paths[new_idx],
                self._audio_paths[idx],
            )
            self._render_audio_list()

    def _select_pptx(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar PowerPoint", "",
            "PowerPoint (*.pptx *.ppt)",
        )
        if path:
            self._pptx_path = path
            self._pptx_label.setText(os.path.basename(path))
            self._pptx_label.setStyleSheet("color: white;")

    def _clear_pptx(self):
        self._pptx_path = ""
        self._pptx_label.setText("Ningún archivo seleccionado")
        self._pptx_label.setStyleSheet("color: gray;")

    def _accept(self):
        if not self._audio_paths:
            QMessageBox.critical(self, "Error", "Agrega al menos un archivo de audio.")
            return
        if not self._asig_combo.currentText().strip():
            QMessageBox.critical(self, "Error", "Ingresa o selecciona la asignatura.")
            return
        self.result_item = QueueItem(
            audio_paths=list(self._audio_paths),
            asignatura=self._asig_combo.currentText().strip(),
            fecha=self._fecha_entry.text().strip(),
            pptx_path=self._pptx_path,
            nombre=self._nombre_entry.text().strip(),
        )
        self.accept()


class ClaseObsidianApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Clase → Obsidian")
        self.setFixedSize(680, 780)
        self.queue: list[QueueItem] = []
        self._selected_row: int | None = None
        self._save_folder: str | None = None
        self._whisper_model = None
        self._model = "large-v3"
        self._fmt = "md"
        self._dest = "obsidian"
        self._setup_ui()

    def _setup_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(8)

        title = QLabel("🎙️ Clase → Obsidian")
        title.setStyleSheet("font-size: 22pt; font-weight: bold;")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title)

        subtitle = QLabel("Transcripción automática de clases universitarias")
        subtitle.setStyleSheet("color: gray; font-size: 10pt;")
        subtitle.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(subtitle)
        layout.addSpacing(4)

        self._tabs = QTabWidget()
        self._tabs.setFixedHeight(330)
        layout.addWidget(self._tabs)

        tab_principal = QWidget()
        tab_opciones = QWidget()
        self._tabs.addTab(tab_principal, "  Principal  ")
        self._tabs.addTab(tab_opciones, "  Opciones  ")

        self._setup_principal_tab(tab_principal)
        self._setup_opciones_tab(tab_opciones)

        self._process_btn = QPushButton("🚀  Procesar Cola")
        self._process_btn.setFixedHeight(46)
        self._process_btn.setStyleSheet(
            "QPushButton { font-size: 14pt; font-weight: bold; }"
            "QPushButton:hover { background-color: #2d5a8a; }"
        )
        self._process_btn.clicked.connect(self._start_processing)
        layout.addWidget(self._process_btn)

        self._progress_label = QLabel("")
        self._progress_label.setStyleSheet("color: gray; font-size: 9pt;")
        layout.addWidget(self._progress_label)

        self._progress_bar = QProgressBar()
        self._progress_bar.setRange(0, 100)
        self._progress_bar.setValue(0)
        self._progress_bar.setFixedHeight(18)
        self._progress_bar.setTextVisible(False)
        layout.addWidget(self._progress_bar)

        status_lbl = QLabel("📋  Estado:")
        status_lbl.setStyleSheet("font-weight: bold;")
        layout.addWidget(status_lbl)

        self._log_box = QTextEdit()
        self._log_box.setReadOnly(True)
        self._log_box.setFixedHeight(150)
        layout.addWidget(self._log_box)

    def _setup_principal_tab(self, tab: QWidget):
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(6)

        toolbar = QHBoxLayout()
        toolbar.setSpacing(6)
        btn_add = QPushButton("+ Agregar")
        btn_add.clicked.connect(self._add_item)
        btn_edit = QPushButton("✏ Editar")
        btn_edit.clicked.connect(self._edit_item)
        btn_rm = QPushButton("✕ Quitar")
        btn_rm.setProperty("class", "gray")
        btn_rm.clicked.connect(self._remove_item)
        btn_apply = QPushButton("↑ Aplicar a todos")
        btn_apply.setProperty("class", "gray")
        btn_apply.clicked.connect(self._apply_to_all)
        toolbar.addWidget(btn_add)
        toolbar.addWidget(btn_edit)
        toolbar.addWidget(btn_rm)
        toolbar.addWidget(btn_apply)
        toolbar.addStretch()
        layout.addLayout(toolbar)

        cols = ["#", "Archivos", "Asignatura", "Fecha", "PPT", "Estado"]
        widths = [35, 160, 140, 90, 40, 90]
        self._table = QTableWidget(0, len(cols))
        self._table.setHorizontalHeaderLabels(cols)
        self._table.verticalHeader().setVisible(False)
        self._table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self._table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self._table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._table.setAlternatingRowColors(False)
        self._table.setShowGrid(False)
        for i, w in enumerate(widths):
            self._table.setColumnWidth(i, w)
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.setRowHeight(0, 28)
        self._table.itemSelectionChanged.connect(self._on_row_select)
        layout.addWidget(self._table)

    def _setup_opciones_tab(self, tab: QWidget):
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(6)

        # Whisper model
        model_lbl = QLabel("Modelo Whisper")
        model_lbl.setStyleSheet("font-weight: bold;")
        layout.addWidget(model_lbl)
        model_row = QHBoxLayout()
        model_row.setSpacing(20)
        self._model_group = QButtonGroup(self)
        for val, text in [("small", "small  (~5 min)"),
                          ("medium", "medium  (~12 min)"),
                          ("large-v3", "large-v3  (~20 min)")]:
            rb = QRadioButton(text)
            if val == self._model:
                rb.setChecked(True)
            rb.toggled.connect(lambda checked, v=val: self._on_model_changed(checked, v))
            self._model_group.addButton(rb)
            model_row.addWidget(rb)
        model_row.addStretch()
        layout.addLayout(model_row)
        layout.addSpacing(8)

        # Format
        fmt_lbl = QLabel("Formato del documento")
        fmt_lbl.setStyleSheet("font-weight: bold;")
        layout.addWidget(fmt_lbl)
        fmt_row = QHBoxLayout()
        fmt_row.setSpacing(20)
        self._fmt_group = QButtonGroup(self)
        for val, text in [("md", "Markdown (.md)"), ("pdf", "PDF (.pdf)")]:
            rb = QRadioButton(text)
            if val == self._fmt:
                rb.setChecked(True)
            rb.toggled.connect(lambda checked, v=val: self._on_fmt_changed(checked, v))
            self._fmt_group.addButton(rb)
            fmt_row.addWidget(rb)
        fmt_row.addStretch()
        layout.addLayout(fmt_row)
        layout.addSpacing(8)

        # Destination
        dest_lbl = QLabel("Destino del documento")
        dest_lbl.setStyleSheet("font-weight: bold;")
        layout.addWidget(dest_lbl)
        dest_row = QHBoxLayout()
        dest_row.setSpacing(20)
        self._dest_group = QButtonGroup(self)
        for val, text in [("obsidian", "Vault de Obsidian"),
                          ("file", "Guardar como archivo...")]:
            rb = QRadioButton(text)
            if val == self._dest:
                rb.setChecked(True)
            rb.toggled.connect(lambda checked, v=val: self._on_dest_changed(checked, v))
            self._dest_group.addButton(rb)
            dest_row.addWidget(rb)
        dest_row.addStretch()
        layout.addLayout(dest_row)
        layout.addSpacing(4)

        # Vault path frame
        self._vault_frame = QWidget()
        vault_layout = QHBoxLayout(self._vault_frame)
        vault_layout.setContentsMargins(0, 0, 0, 0)
        vault_layout.setSpacing(8)
        self._vault_entry = QLineEdit()
        self._vault_entry.setText(os.getenv("OBSIDIAN_VAULT", r"C:\Users\vicen\Desktop\claude\Claude"))
        btn_vault = QPushButton("Cambiar")
        btn_vault.setFixedWidth(80)
        btn_vault.clicked.connect(self._select_vault)
        vault_layout.addWidget(self._vault_entry)
        vault_layout.addWidget(btn_vault)
        layout.addWidget(self._vault_frame)
        layout.addStretch()

        self._vault_frame.setVisible(self._dest == "obsidian")

    # ── option handlers ───────────────────────────────────────────────────────

    def _on_model_changed(self, checked: bool, val: str):
        if checked:
            self._model = val

    def _on_fmt_changed(self, checked: bool, val: str):
        if checked:
            self._fmt = val

    def _on_dest_changed(self, checked: bool, val: str):
        if checked:
            self._dest = val
            self._vault_frame.setVisible(val == "obsidian")

    def _select_vault(self):
        path = QFileDialog.getExistingDirectory(self, "Seleccionar vault de Obsidian")
        if path:
            self._vault_entry.setText(path)

    # ── queue table ───────────────────────────────────────────────────────────

    def _audio_display(self, item: QueueItem) -> str:
        if not item.audio_paths:
            return "—"
        name = os.path.basename(item.audio_paths[0])
        if len(name) > 20:
            name = name[:17] + "..."
        if len(item.audio_paths) > 1:
            name += f" (+{len(item.audio_paths) - 1})"
        return name

    def _refresh_table(self):
        estado_colors = {
            "pendiente": "#ffffff",
            "procesando": "#f6e05e",
            "listo": "#68d391",
            "error": "#fc8181",
        }
        self._table.setRowCount(len(self.queue))
        for i, item in enumerate(self.queue):
            asig = item.asignatura
            if len(asig) > 18:
                asig = asig[:15] + "..."
            values = [
                str(i + 1),
                self._audio_display(item),
                asig or "—",
                item.fecha or "—",
                "✓" if item.pptx_path else "—",
                item.estado,
            ]
            color = QColor(estado_colors.get(item.estado, "#ffffff"))
            for j, val in enumerate(values):
                cell = QTableWidgetItem(val)
                cell.setForeground(color)
                self._table.setItem(i, j, cell)
                self._table.setRowHeight(i, 28)

        if self._selected_row is not None and self._selected_row < len(self.queue):
            self._table.selectRow(self._selected_row)

    def _on_row_select(self):
        rows = self._table.selectedItems()
        if rows:
            self._selected_row = self._table.currentRow()

    # ── queue actions ─────────────────────────────────────────────────────────

    def _add_item(self):
        vault = self._vault_entry.text().strip()
        dlg = EditQueueItemDialog(self, vault_folders=scan_vault_folders(vault))
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.result_item:
            self.queue.append(dlg.result_item)
            self._selected_row = len(self.queue) - 1
            self._refresh_table()

    def _edit_item(self):
        if self._selected_row is None or self._selected_row >= len(self.queue):
            QMessageBox.warning(self, "Advertencia", "Selecciona un elemento de la cola.")
            return
        item = self.queue[self._selected_row]
        if item.estado == "procesando":
            QMessageBox.warning(self, "Advertencia", "No se puede editar un elemento en proceso.")
            return
        vault = self._vault_entry.text().strip()
        dlg = EditQueueItemDialog(self, item=item, vault_folders=scan_vault_folders(vault))
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.result_item:
            dlg.result_item.estado = item.estado
            self.queue[self._selected_row] = dlg.result_item
            self._refresh_table()

    def _remove_item(self):
        if self._selected_row is None or self._selected_row >= len(self.queue):
            QMessageBox.warning(self, "Advertencia", "Selecciona un elemento de la cola.")
            return
        if self.queue[self._selected_row].estado == "procesando":
            QMessageBox.warning(self, "Advertencia", "No se puede quitar un elemento en proceso.")
            return
        self.queue.pop(self._selected_row)
        if self.queue:
            self._selected_row = min(self._selected_row, len(self.queue) - 1)
        else:
            self._selected_row = None
        self._refresh_table()

    def _apply_to_all(self):
        if self._selected_row is None or self._selected_row >= len(self.queue):
            QMessageBox.warning(self, "Advertencia", "Selecciona el elemento fuente.")
            return
        src = self.queue[self._selected_row]
        for i, item in enumerate(self.queue):
            if i != self._selected_row:
                item.asignatura = src.asignatura
                item.fecha = src.fecha
        self._refresh_table()

    # ── processing ────────────────────────────────────────────────────────────

    def _start_processing(self):
        pending = [i for i, item in enumerate(self.queue) if item.estado == "pendiente"]
        if not pending:
            QMessageBox.critical(self, "Error", "No hay elementos pendientes en la cola.")
            return

        if self._dest == "file":
            folder = QFileDialog.getExistingDirectory(self, "Seleccionar carpeta de destino")
            if not folder:
                return
            self._save_folder = folder
        else:
            self._save_folder = None

        self._process_btn.setEnabled(False)
        self._process_btn.setText("⏳  Procesando...")
        self._log_box.clear()
        self._progress_bar.setValue(0)
        self._progress_label.setText("")

        self._worker = _Worker(self)
        self._worker.log_signal.connect(self._append_log)
        self._worker.progress_signal.connect(self._set_progress)
        self._worker.refresh_signal.connect(self._refresh_table)
        self._worker.done_signal.connect(self._on_done)
        self._worker.start()

    def _on_done(self):
        self._process_btn.setEnabled(True)
        self._process_btn.setText("🚀  Procesar Cola")

    def _append_log(self, msg: str):
        self._log_box.append(msg)
        cursor = self._log_box.textCursor()
        cursor.movePosition(QTextCursor.MoveOperation.End)
        self._log_box.setTextCursor(cursor)

    def _set_progress(self, value: float, label: str = ""):
        self._progress_bar.setValue(int(value * 100))
        self._progress_label.setText(label)

    def _save(self, content: str, asignatura: str, fecha: str, nombre: str = "") -> str:
        ext = ".pdf" if self._fmt == "pdf" else ".md"
        base = nombre if nombre else f"{fecha} - {asignatura}"
        filename = f"{base}{ext}"

        if self._dest == "obsidian":
            vault = self._vault_entry.text().strip()
            subfolder = os.path.join(vault, asignatura)
            folder = subfolder if os.path.isdir(subfolder) else vault
            filepath = os.path.join(folder, filename)
        else:
            filepath = os.path.join(self._save_folder, filename)

        if self._fmt == "pdf":
            save_pdf(content, filepath)
        else:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)

        return filepath


class SetupDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Configuración inicial")
        self.setFixedSize(460, 280)
        self.api_key = None
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(30, 24, 30, 20)
        layout.setSpacing(8)

        title = QLabel("🎙️ Clase → Obsidian")
        title.setStyleSheet("font-size: 18pt; font-weight: bold;")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title)

        sub = QLabel("Configuración inicial — solo se hace una vez")
        sub.setStyleSheet("color: gray;")
        sub.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(sub)
        layout.addSpacing(12)

        key_lbl = QLabel("API Key de Anthropic:")
        key_lbl.setStyleSheet("font-weight: bold;")
        layout.addWidget(key_lbl)

        self._key_entry = QLineEdit()
        self._key_entry.setPlaceholderText("sk-ant-api03-...")
        self._key_entry.setEchoMode(QLineEdit.EchoMode.Password)
        self._key_entry.setFixedHeight(38)
        layout.addWidget(self._key_entry)

        btn_toggle = QPushButton("👁 Mostrar / Ocultar")
        btn_toggle.setProperty("class", "gray")
        btn_toggle.setFixedWidth(160)
        btn_toggle.clicked.connect(self._toggle_show)
        layout.addWidget(btn_toggle, alignment=Qt.AlignmentFlag.AlignLeft)
        layout.addSpacing(8)

        btn_save = QPushButton("Guardar y continuar")
        btn_save.setFixedHeight(42)
        btn_save.setStyleSheet("font-size: 13pt; font-weight: bold;")
        btn_save.clicked.connect(self._save)
        layout.addWidget(btn_save)

    def _toggle_show(self):
        if self._key_entry.echoMode() == QLineEdit.EchoMode.Password:
            self._key_entry.setEchoMode(QLineEdit.EchoMode.Normal)
        else:
            self._key_entry.setEchoMode(QLineEdit.EchoMode.Password)

    def _save(self):
        key = self._key_entry.text().strip()
        if not key.startswith("sk-ant-"):
            QMessageBox.critical(self, "Error", "La key debe comenzar con 'sk-ant-'")
            return
        env_path = os.path.join(_base, ".env")
        vault = os.getenv("OBSIDIAN_VAULT", r"C:\Users\vicen\Desktop\claude\Claude")
        with open(env_path, "w", encoding="utf-8") as f:
            f.write(f"ANTHROPIC_API_KEY={key}\n")
            f.write(f"OBSIDIAN_VAULT={vault}\n")
        load_dotenv(dotenv_path=env_path, override=True)
        self.api_key = key
        self.accept()

    def closeEvent(self, event):
        if not self.api_key:
            reply = QMessageBox.question(
                self, "Salir",
                "Sin API key la app no puede funcionar. ¿Salir?",
                QMessageBox.StandardButton.Ok | QMessageBox.StandardButton.Cancel,
            )
            if reply == QMessageBox.StandardButton.Ok:
                event.accept()
                sys.exit(0)
            else:
                event.ignore()
        else:
            event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setStyleSheet(DARK_STYLESHEET)

    if not os.getenv("ANTHROPIC_API_KEY"):
        setup = SetupDialog()
        if setup.exec() != QDialog.DialogCode.Accepted:
            sys.exit(0)

    window = ClaseObsidianApp()
    window.show()
    sys.exit(app.exec())
